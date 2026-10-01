"""
W0-03E-R1 / C01 — tenant-scoped offer and finance exports.

The final W0-03E/C03 review reproduced a cross-tenant leak through the actual
HTTP response: ``GET /api/offers/{id}/xlsx`` loaded tenant A's offer, then the
related ``projects`` document by the bare ``offer.project_id`` — tenant B's
project with the same id came back in cell B3. The PDF exporter had the same
bare join. These regressions call the real routes (TestClient, FastAPI app)
over one shared legacy database in which tenant B holds:

  * a project with the SAME id as A's project (inserted first, so a bare
    ``find_one`` returns B's);
  * a project that exists ONLY in B, which one of A's offers / invoices names;
  * an offer and offer events with A's offer id;
  * counterparties and users with A's ids (from the C03 world).

Every response — XLSX cells, PDF text, JSON — must contain zero B names,
codes, values or canonical ids, in ``off``, ``shadow`` and ``enforce``. A
reference that resolves only in B yields no project (empty / "Unknown"),
never an inferred one. Same-tenant behaviour is unchanged (A's own project
appears). No route writes anything into B.

Run:  pytest tests/test_w0_03e_r1_exports.py -v --noconftest
"""
import base64
import io
import re
import zlib

import pytest

from app.routes import dashboard, finance, offers

from tests.test_w0_03e_c03_isolation import (
    B_MARKERS, MODES, _world, b_canonical_ids, set_mode,
)
from tests.test_w0_03e_legacy_migration import ORG_A, ORG_B, run
from tests.test_w0_03e_legacy_routes import ADMIN, client_for

#: B-only strings added by this module, on top of the C03 markers.
R1_MARKERS = B_MARKERS + ("B-SECRET-PROJECT", "B-CODE-PRX", "B-CODE-ONLY", "B-SECRET-OFFER",
                          "B-SECRET-EVENT", "B-SECRET-ADDRESS")
A_PROJECT = "A-CODE-PRX - A-OWN-PROJECT"
PERIOD = "date_from=2026-01-01&date_to=2026-12-31"


def _line(name):
    return {"id": "l-" + name, "activity_name": name, "unit": "m2", "qty": 2.0,
            "material_unit_cost": 1.5, "labor_unit_cost": 2.5, "note": ""}


async def _r1_world(db=None):
    db = await _world(db)
    # tenant B's project with A's id (B's copy was inserted first by the C03 world)
    await db["projects"].update_one({"id": "pr-x", "org_id": ORG_B}, {"$set": {
        "code": "B-CODE-PRX", "name": "B-SECRET-PROJECT", "address_text": "B-SECRET-ADDRESS"}})
    await db["projects"].update_one({"id": "pr-x", "org_id": ORG_A}, {"$set": {
        "code": "A-CODE-PRX", "name": "A-OWN-PROJECT", "address_text": "A street"}})
    await db["projects"].insert_one({"id": "pr-only-b", "org_id": ORG_B, "code": "B-CODE-ONLY",
                                     "name": "B-SECRET-PROJECT", "address_text": "B-SECRET-ADDRESS"})
    # tenant B: an offer and its events under A's offer id
    await db["offers"].insert_one({
        "id": "offer-a", "org_id": ORG_B, "project_id": "pr-x", "offer_no": "B-SECRET-OFFER",
        "status": "Sent", "review_token": "tok-b", "created_at": "2026-09-01T00:00:00",
        "lines": [_line("B-SECRET-OFFER line")], "currency": "BGN", "subtotal": 999999,
        "vat_percent": 0, "vat_amount": 0, "total": 999999})
    await db["offer_events"].insert_one({
        "id": "ev-b", "org_id": ORG_B, "offer_id": "offer-a", "event_type": "viewed",
        "actor": "B-SECRET-EVENT", "created_at": "2026-09-01", "details": {}})
    # tenant A: one offer on its own (shared-id) project, one naming a B-only project
    for oid, pid, tok in (("offer-a", "pr-x", "tok-a"), ("offer-a-foreign", "pr-only-b", "tok-af")):
        await db["offers"].insert_one({
            "id": oid, "org_id": ORG_A, "project_id": pid, "offer_no": "OFF-%s" % oid,
            "title": "Оферта А", "status": "Sent", "review_token": tok, "version": 1,
            "created_at": "2026-09-01T00:00:00", "lines": [_line("Мазилка")], "currency": "BGN",
            "subtotal": 8.0, "vat_percent": 20, "vat_amount": 1.6, "total": 9.6})
    # invoices: A's point at the shared-id project and at the B-only one
    await db["invoices"].update_one({"id": "inv-shared", "org_id": ORG_A}, {"$set": {
        "project_id": "pr-x", "allocations": [{"type": "project", "ref_id": "pr-x"}]}})
    await db["invoices"].update_one({"id": "inv-a-foreign", "org_id": ORG_A}, {"$set": {
        "project_id": "pr-only-b", "allocations": [{"type": "project", "ref_id": "pr-only-b"}]}})
    await db["invoices"].update_many({"org_id": ORG_B, "id": {"$in": ["inv-shared", "inv-only-b"]}},
                                     {"$set": {"project_id": "pr-x",
                                               "allocations": [{"type": "project", "ref_id": "pr-x"}]}})
    return db


def r1_world():
    return run(_r1_world())


def assert_no_b(db, text, b_ids=None, echoed=()):
    text = text if isinstance(text, str) else repr(text)
    for marker in R1_MARKERS:
        if marker not in echoed:
            assert marker not in text, "tenant B value %r reached tenant A" % marker
    ids = b_ids if b_ids is not None else b_canonical_ids(db)
    leaked = [i for i in ids if i in text]
    assert not leaked, "tenant B canonical ids reached tenant A: %s" % leaked[:3]


def xlsx_text(content):
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(content))
    return {c.coordinate: c.value for ws in wb for row in ws.iter_rows() for c in row
            if c.value is not None}


def pdf_text(content):
    """Every content stream of the PDF, inflated. The route falls back to the
    standard Helvetica font in these tests (see ``plain_pdf_fonts``), so ASCII
    text is written literally into the stream and can be searched."""
    out = []
    for head, raw in re.findall(rb"<<((?:(?!<<).)*?)>>\s*stream\r?\n(.*?)\r?\n?endstream", content, re.S):
        data = raw.strip()
        if b"ASCII85Decode" in head:
            data = base64.a85decode(data[:-2] if data.endswith(b"~>") else data)
        if b"FlateDecode" in head:
            data = zlib.decompress(data)
        out.append(data)
    text = b"\n".join(out).decode("latin-1")
    assert text.strip(), "no PDF content stream could be read"
    return text


@pytest.fixture
def plain_pdf_fonts(monkeypatch):
    """Make the DejaVu TrueType font unavailable, so the exporter takes its
    existing Helvetica fallback (subset TTF glyph ids are not searchable text)."""
    import reportlab.pdfbase.ttfonts as ttfonts

    def unavailable(*a, **kw):
        raise OSError("font unavailable in this test")
    monkeypatch.setattr(ttfonts, "TTFont", unavailable)


def _snapshot_b(db):
    async def snap():
        out = {}
        for name in ("projects", "offers", "offer_events", "invoices", "counterparties"):
            out[name] = await db[name].find({"org_id": ORG_B}, {"_id": 0}).sort("id", 1).to_list(None)
        return out
    return run(snap())


# ================================================================ offer XLSX — the reproduced leak
@pytest.mark.parametrize("mode", MODES)
def test_offer_xlsx_reads_the_project_only_inside_the_callers_tenant(monkeypatch, mode):
    db = r1_world()
    before = _snapshot_b(db)
    set_mode(monkeypatch, db, mode)
    c = client_for(monkeypatch, offers, db, ADMIN)
    r = c.get("/api/offers/offer-a/xlsx")
    assert r.status_code == 200, r.text
    cells = xlsx_text(r.content)
    assert cells["B3"] == A_PROJECT                       # A's own pr-x, not B's pr-x
    assert "OFF-offer-a" in cells["A2"]                   # A's offer, not B's same-id offer
    assert_no_b(db, " ".join(map(str, cells.values())))
    r = c.get("/api/offers/offer-a-foreign/xlsx")         # names a project that exists only in B
    assert r.status_code == 200, r.text
    cells = xlsx_text(r.content)
    assert "B3" not in cells                              # no project — never B's
    assert_no_b(db, " ".join(map(str, cells.values())))
    assert _snapshot_b(db) == before                      # nothing of B changed


# ================================================================ offer PDF — the parallel bare join
@pytest.mark.parametrize("mode", MODES)
def test_offer_pdf_reads_the_project_only_inside_the_callers_tenant(monkeypatch, plain_pdf_fonts,
                                                                    mode):
    db = r1_world()
    before = _snapshot_b(db)
    set_mode(monkeypatch, db, mode)
    c = client_for(monkeypatch, offers, db, ADMIN)
    r = c.get("/api/offers/offer-a/pdf")
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == "application/pdf"
    text = pdf_text(r.content)
    assert "A-CODE-PRX" in text and "A-OWN-PROJECT" in text   # extraction works: A's project
    assert "OFF-offer-a" in text
    assert_no_b(db, text)
    r = c.get("/api/offers/offer-a-foreign/pdf")
    assert r.status_code == 200, r.text
    text = pdf_text(r.content)
    assert "OFF-offer-a-foreign" in text                      # the offer itself is exported
    assert "A-OWN-PROJECT" not in text                        # no project is inferred
    assert_no_b(db, text)
    assert _snapshot_b(db) == before


def test_pdf_text_extraction_would_see_a_b_project(plain_pdf_fonts, monkeypatch):
    """Control for the PDF assertions above: had the route read B's project, the
    extracted text WOULD contain its code — so 'absent' is a real result."""
    db = r1_world()
    run(db["projects"].update_one({"id": "pr-x", "org_id": ORG_A}, {"$set": {
        "code": "B-CODE-PRX", "name": "B-SECRET-PROJECT"}}))   # A's own row, B-looking text
    monkeypatch.setenv("MASTER_DATA_MODE", "off")
    text = pdf_text(client_for(monkeypatch, offers, db, ADMIN).get("/api/offers/offer-a/pdf").content)
    assert "B-CODE-PRX" in text and "B-SECRET-PROJECT" in text


def test_offer_exports_of_a_foreign_offer_id_are_not_found(monkeypatch, plain_pdf_fonts):
    db = r1_world()
    run(db["offers"].insert_one({"id": "offer-only-b", "org_id": ORG_B, "project_id": "pr-only-b",
                                 "offer_no": "B-SECRET-OFFER", "lines": []}))
    monkeypatch.setenv("MASTER_DATA_MODE", "off")
    c = client_for(monkeypatch, offers, db, ADMIN)
    for fmt in ("xlsx", "pdf"):
        r = c.get("/api/offers/offer-only-b/%s" % fmt)
        assert r.status_code == 404
        assert_no_b(db, r.text, echoed=())


# ================================================================ offer read projections
@pytest.mark.parametrize("mode", MODES)
def test_offer_list_detail_events_and_public_review_show_only_tenant_a(monkeypatch, mode):
    db = r1_world()
    set_mode(monkeypatch, db, mode)
    c = client_for(monkeypatch, offers, db, ADMIN)
    listed = {o["id"]: o for o in c.get("/api/offers").json()}
    assert set(listed) == {"offer-a", "offer-a-foreign"}
    assert (listed["offer-a"]["project_code"], listed["offer-a"]["project_name"]) == (
        "A-CODE-PRX", "A-OWN-PROJECT")
    assert (listed["offer-a-foreign"]["project_code"],
            listed["offer-a-foreign"]["project_name"]) == ("", "")
    detail = c.get("/api/offers/offer-a").json()
    assert detail["project_name"] == "A-OWN-PROJECT" and detail["offer_no"] == "OFF-offer-a"
    foreign = c.get("/api/offers/offer-a-foreign").json()
    assert foreign["project_name"] == ""
    events = c.get("/api/offers/offer-a/events").json()
    assert all(e["org_id"] == ORG_A for e in events)
    for payload in (listed, detail, foreign, events):
        assert_no_b(db, payload)

    # public review: no session — the tenant is the org of the offer the token names
    review = c.get("/api/offers/review/tok-a").json()
    assert review["project_name"] == "A-OWN-PROJECT" and review["project_address"] == "A street"
    assert c.get("/api/offers/review/tok-af").json()["project_name"] == ""
    assert_no_b(db, review)
    # A's view is recorded in A even though B has a 'viewed' event for the same offer id
    a_views = run(db["offer_events"].count_documents(
        {"org_id": ORG_A, "offer_id": "offer-a", "event_type": "viewed"}))
    b_views = run(db["offer_events"].count_documents(
        {"org_id": ORG_B, "offer_id": "offer-a", "event_type": "viewed"}))
    assert (a_views, b_views) == (1, 1)
    events = c.get("/api/offers/offer-a/events").json()
    assert [e["event_type"] for e in events] == ["viewed"] and events[0]["org_id"] == ORG_A
    assert_no_b(db, events)


# ================================================================ client-invoice PDF
@pytest.mark.parametrize("mode", MODES)
def test_client_invoice_pdf_reads_the_project_only_inside_the_callers_tenant(
        monkeypatch, plain_pdf_fonts, mode):
    db = r1_world()
    before = _snapshot_b(db)
    set_mode(monkeypatch, db, mode)
    c = client_for(monkeypatch, finance, db, ADMIN)
    r = c.get("/api/finance/invoices/inv-shared/pdf")
    assert r.status_code == 200, r.text
    text = pdf_text(r.content)
    assert "A-INV-1" in text                                 # A's invoice, not B's same-id one
    assert "A-OWN-PROJECT" in text
    assert_no_b(db, text)
    r = c.get("/api/finance/invoices/inv-a-foreign/pdf")      # project exists only in B
    assert r.status_code == 200, r.text
    text = pdf_text(r.content)
    assert "A-INV-2" in text and "A-OWN-PROJECT" not in text
    assert_no_b(db, text)
    assert c.get("/api/finance/invoices/inv-only-b/pdf").status_code == 404
    assert _snapshot_b(db) == before


# ================================================================ finance drill-downs
@pytest.mark.parametrize("mode", MODES)
def test_finance_detail_drilldowns_name_only_tenant_a_records(monkeypatch, mode):
    db = r1_world()
    set_mode(monkeypatch, db, mode)
    c = client_for(monkeypatch, dashboard, db, ADMIN)
    by_cp = c.get("/api/reports/finance-details/by-counterparty?" + PERIOD).json()
    names = {r["counterparty_id"]: r["counterparty_name"] for r in by_cp["items"]}
    assert names == {"cp1": "Baumit Bulgaria", "cp-only-b": "Unknown"}
    assert sum(r["total"] for r in by_cp["items"]) == 15       # A's 10 + 5; B's 999999 never
    by_project = c.get("/api/reports/finance-details/by-project?" + PERIOD).json()
    projects = {r["project_id"]: (r["project_code"], r["project_name"]) for r in by_project["items"]}
    assert projects == {"pr-x": ("A-CODE-PRX", "A-OWN-PROJECT"), "pr-only-b": ("", "Unknown")}
    top = {}
    for direction in ("expense", "income"):
        body = c.get("/api/reports/finance-details/top-counterparties?direction=%s&%s"
                     % (direction, PERIOD)).json()
        top[direction] = {r["counterparty_id"]: r["counterparty_name"] for r in body["items"]}
        assert_no_b(db, body)
    assert top == {"expense": {"cp1": "Baumit Bulgaria"}, "income": {"cp-only-b": "Unknown"}}
    tx = c.get("/api/reports/finance-details/transactions?transaction_type=invoice&" + PERIOD).json()
    assert_no_b(db, tx)
    for body in (by_cp, by_project):
        assert_no_b(db, body)


def test_finance_drilldowns_keep_same_tenant_names_when_b_has_none(monkeypatch):
    """Same-tenant behaviour preserved: with no B data at all the answers are
    identical to the scoped answers above."""
    from tests.test_w0_03e_legacy_migration import new_db
    db = r1_world()

    async def only_a():
        clean = new_db()
        for name in ("invoices", "counterparties", "projects"):
            for d in await db[name].find({"org_id": ORG_A}, {"_id": 0}).to_list(None):
                await clean[name].insert_one(d)
        return clean
    clean = run(only_a())
    monkeypatch.setenv("MASTER_DATA_MODE", "off")
    scoped = client_for(monkeypatch, dashboard, db, ADMIN).get(
        "/api/reports/finance-details/by-project?" + PERIOD).json()["items"]
    alone = client_for(monkeypatch, dashboard, clean, ADMIN).get(
        "/api/reports/finance-details/by-project?" + PERIOD).json()["items"]
    assert sorted(scoped, key=lambda r: r["project_id"]) == sorted(alone, key=lambda r: r["project_id"])
