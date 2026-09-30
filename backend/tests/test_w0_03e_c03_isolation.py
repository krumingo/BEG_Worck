"""
W0-03E / C03 — final bounded correction: tenant isolation of the assigned
report/export/import/adapter surface.

The C02 re-review reproduced a leak on a real server: ``/prices`` matched the
caller's ``invoice_lines`` but joined ``invoices`` by the bare ``invoice_id``,
so a tenant-B invoice sharing that id returned B's invoice number and supplier
inside tenant A's response. These regressions prove, for every scenario the C03
assignment names, that nothing of tenant B — name, value, legacy id that only B
holds, canonical Master id or audit — appears in what tenant A receives.

Every scenario runs in ``off``, ``shadow`` and ``enforce`` where the path has
mode-dependent behaviour. A single shared legacy database holds both orgs (the
hardest case: database-per-tenant separates them physically anyway); both
tenants are migrated so that B also HAS canonical Master ids that could leak.

A value that tenant A itself stored — for example the id string of a record
that exists only in B, written into A's own invoice — is A's data and may be
echoed back as A stored it. What must never come back is anything read FROM B.

Run:  pytest tests/test_w0_03e_c03_isolation.py -v --noconftest
"""
import io
import json

import pytest

from app.audit.store import AUDIT_COLLECTION
from app.master_data import legacy_adapter as la
from app.master_data import legacy_plan as lp
from app.master_data.deps import ENV_MODE
from app.master_data.models import build_entity, new_legacy_ref
from app.routes import hr, projects, reports

from tests.test_w0_03e_legacy_migration import (
    ORG_A, ORG_B, T_A, T_B, Ctx, dry, go, repo, run, seed, seeded,
)
from tests.test_w0_03e_legacy_routes import ADMIN, client_for, enforce_with, find_all, landmines

MODES = ["off", "shadow", "enforce"]

#: Strings that exist ONLY in tenant B's documents. None may reach tenant A.
B_MARKERS = ("B-SECRET-INVOICE", "B-SECRET-SUPPLIER", "B-SECRET-USER", "B-SECRET-WH",
             "B-SECRET-COMPANY", "B-SECRET-CLIENT", "999999", "B-ONLY-EIK-777",
             "cp-b-hidden-supplier")


async def _world(db=None):
    """The A/B world. ``db`` given: build it there (the real-Mongo tests)."""
    if db is None:
        db = await seeded()                       # tenant A's legacy data + one existing Master
    else:
        from tests.test_w0_03e_real_mongo import prepare
        await prepare(db)
    await seed(db, ORG_B)                         # tenant B: the SAME legacy ids, same database
    # B documents that share ids with A's documents, with B-only content
    b_docs = {
        "invoices": [
            {"id": "inv-shared", "org_id": ORG_B, "invoice_no": "B-SECRET-INVOICE",
             "issue_date": "2026-09-02", "supplier_counterparty_id": "cp-b-hidden-supplier",
             "direction": "Received", "total": 999999, "subtotal": 999999, "vat_amount": 0,
             "paid_amount": 0, "remaining_amount": 999999},
            {"id": "inv-only-b", "org_id": ORG_B, "invoice_no": "B-SECRET-INVOICE",
             "issue_date": "2026-09-02", "supplier_counterparty_id": "cp1",
             "direction": "Issued", "total": 999999, "subtotal": 999999, "vat_amount": 0,
             "paid_amount": 0, "remaining_amount": 999999}],
        "counterparties": [
            {"id": "cp-b-hidden-supplier", "org_id": ORG_B, "name": "B-SECRET-SUPPLIER",
             "eik": "B-ONLY-EIK-777", "type": "client"},
            {"id": "cp-only-b", "org_id": ORG_B, "name": "B-SECRET-SUPPLIER", "type": "client",
             "eik": "B-ONLY-EIK-777"}],
        "users": [{"id": "u-only-b", "org_id": ORG_B, "first_name": "B-SECRET-USER",
                   "last_name": "X", "name": "B-SECRET-USER"}],
        "warehouses": [{"id": "wh-only-b", "org_id": ORG_B, "code": "B-SECRET-WH"}],
        "companies": [{"id": "own-x", "org_id": ORG_B, "name": "B-SECRET-COMPANY",
                       "eik": "B-ONLY-EIK-777"}],
        "projects": [{"id": "pr-x", "org_id": ORG_B, "owner_type": "company", "owner_id": "own-x"}],
    }
    for name, docs in b_docs.items():
        for d in docs:
            await db[name].insert_one(dict(d))
    # B's copies of shared-id records carry B-only names
    await db["users"].update_one({"id": "u1", "org_id": ORG_B}, {"$set": {"first_name": "B-SECRET-USER"}})
    await db["warehouses"].update_one({"id": "wh1", "org_id": ORG_B}, {"$set": {"code": "B-SECRET-WH"}})
    await db["counterparties"].update_one({"id": "cp1", "org_id": ORG_B},
                                          {"$set": {"name": "B-SECRET-SUPPLIER"}})
    await db["companies"].update_one({"id": "c1", "org_id": ORG_B},
                                     {"$set": {"name": "B-SECRET-COMPANY", "eik": "B-ONLY-EIK-777"}})
    # tenant A: an invoice sharing B's id, one pointing at an invoice only B has,
    # and rows that reference ids present only in B (as an import would store them)
    a_docs = {
        "invoices": [
            {"id": "inv-shared", "org_id": ORG_A, "invoice_no": "A-INV-1", "issue_date": "2026-09-01",
             "supplier_counterparty_id": "cp1", "direction": "Received", "total": 10, "subtotal": 8,
             "vat_amount": 2, "paid_amount": 0, "remaining_amount": 10},
            {"id": "inv-a-foreign", "org_id": ORG_A, "invoice_no": "A-INV-2",
             "issue_date": "2026-09-01", "supplier_counterparty_id": "cp-only-b",
             "direction": "Issued", "total": 5, "subtotal": 4, "vat_amount": 1, "paid_amount": 0,
             "remaining_amount": 5}],
        "invoice_lines": [
            {"id": "line-a", "org_id": ORG_A, "invoice_id": "inv-shared", "description": "Цимент",
             "purchased_by_user_id": "u1", "created_at": "2026-09-01",
             "allocations": [{"type": "warehouse", "ref_id": "wh1", "qty": 1},
                             {"type": "warehouse", "ref_id": "wh-only-b", "qty": 2}]},
            {"id": "line-a2", "org_id": ORG_A, "invoice_id": "inv-only-b", "description": "Пясък",
             "purchased_by_user_id": "u-only-b", "created_at": "2026-09-02", "allocations": []}],
        "clients": [{"id": "own-x", "org_id": ORG_A, "companyName": "А Клиент ООД", "eik": "121212121"}],
        "projects": [{"id": "pr-x", "org_id": ORG_A, "owner_type": "company", "owner_id": "own-x"}],
        "advances": [{"id": "adv-foreign-user", "org_id": ORG_A, "user_id": "u-only-b", "amount": 1,
                      "issued_date": "2026-09-01"}],
    }
    for name, docs in a_docs.items():
        for d in docs:
            await db[name].insert_one(dict(d))
    # both tenants migrated: B now also has canonical ids that must not leak
    await go(db, (await dry(db))["plan_token"])
    b_ctx = Ctx(db, tenant_id=T_B, org_id=ORG_B, user_id="owner-b")
    await go(db, (await dry(db, ctx=b_ctx))["plan_token"], ctx=b_ctx)
    return db


def world():
    return run(_world())


async def b_canonical_ids_async(db):
    ids = set()
    for etype in ("person", "organization", "activity", "item", "asset_type", "physical_asset",
                  "location"):
        for d in await db["md_" + etype].find({"tenant_id": T_B}, {"_id": 0, "id": 1}).to_list(None):
            ids.add(d["id"])
    return ids


def b_canonical_ids(db):
    return run(b_canonical_ids_async(db))


def assert_no_b(db, payload, echoed=(), b_ids=None):
    """``echoed``: values the CALLER itself put in its request (a path id) and
    that the route may repeat — the caller's input, not data read from B."""
    text = payload if isinstance(payload, str) else json.dumps(payload, ensure_ascii=False)
    for marker in B_MARKERS:
        if marker in echoed:
            continue
        assert marker not in text, "tenant B value %r reached tenant A" % marker
    leaked = [i for i in (b_ids if b_ids is not None else b_canonical_ids(db)) if i in text]
    assert not leaked, "tenant B canonical ids reached tenant A: %s" % leaked[:3]


def set_mode(monkeypatch, db, mode):
    if mode == "enforce":
        return enforce_with(monkeypatch, db)
    monkeypatch.setenv(ENV_MODE, mode)
    return None


# ================================================================ /prices — the reproduced leak
@pytest.mark.parametrize("mode", MODES)
def test_prices_never_joins_another_tenants_invoice(monkeypatch, mode):
    db = world()
    set_mode(monkeypatch, db, mode)
    body = client_for(monkeypatch, reports, db, ADMIN).get("/api/prices").json()
    rows = {r["line_id"]: r for r in body["items"]}
    assert set(rows) == {"line-a", "line-a2"}                 # A's own lines, all of them
    shared = rows["line-a"]
    assert shared["invoice_no"] == "A-INV-1" and shared["supplier_id"] == "cp1"
    assert shared["supplier_name"] == "Baumit Bulgaria"       # A's cp1, not B's cp1
    assert shared["purchased_by_name"] == "Иван Петров"       # A's u1, not B's u1
    assert shared["allocation_summary"] == "W:C1:1, W:?:2"    # A's wh1; B-only wh unknown
    only_b = rows["line-a2"]                                   # its invoice exists only in B
    assert "invoice_no" not in only_b and "supplier_id" not in only_b
    assert only_b["purchased_by_name"] == ""                  # user exists only in B
    assert_no_b(db, body)
    if mode == "enforce":
        assert shared["supplier_master_ref"]["status"] == "mapped"
        assert shared["supplier_master_ref"]["canonical_id"] not in b_canonical_ids(db)
    else:
        assert "supplier_master_ref" not in shared


@pytest.mark.parametrize("mode", MODES)
def test_prices_filters_apply_to_this_tenants_invoice_only(monkeypatch, mode):
    db = world()
    set_mode(monkeypatch, db, mode)
    c = client_for(monkeypatch, reports, db, ADMIN)
    # B's invoice with the shared id names cp-b-hidden-supplier: filtering by it finds nothing
    body = c.get("/api/prices?supplier_id=cp-b-hidden-supplier").json()
    assert body["items"] == [] and body["total"] == 0
    body = c.get("/api/prices?date_from=2026-09-02").json()   # B's invoice date
    assert body["items"] == []
    assert_no_b(db, c.get("/api/prices?sort_by=invoice_date").json())


# ================================================================ turnover, drilldown, client
@pytest.mark.parametrize("mode", MODES)
def test_turnover_reports_and_drilldowns_contain_only_tenant_a(monkeypatch, mode):
    db = world()
    set_mode(monkeypatch, db, mode)
    c = client_for(monkeypatch, reports, db, ADMIN)
    purchases = c.get("/api/reports/turnover-by-counterparty?type=purchases").json()
    assert {r["counterparty_id"]: r["counterparty_name"] for r in purchases["items"]} == {
        "cp1": "Baumit Bulgaria"}
    assert purchases["grand_totals"]["total_amount"] == 10          # B's 999999 not aggregated
    sales = c.get("/api/reports/turnover-by-counterparty?type=sales").json()
    foreign = [r for r in sales["items"] if r["counterparty_id"] == "cp-only-b"][0]
    assert foreign["counterparty_name"] == "(Неизвестен)" and foreign["counterparty_eik"] is None
    clients = c.get("/api/reports/turnover-by-client").json()
    assert [r["client_id"] for r in clients["items"]] == []        # cp-only-b is B's client
    for cp in ("cp1", "cp-only-b", "cp-b-hidden-supplier"):
        drill = c.get("/api/reports/turnover-by-counterparty/%s/invoices?type=all" % cp).json()
        assert all(i["org_id"] == ORG_A for i in drill["items"])
        if cp != "cp1":
            assert drill["counterparty_name"] == "(Неизвестен)"
        assert_no_b(db, drill, echoed=(cp,))
    for body in (purchases, sales, clients):
        assert_no_b(db, body)
    if mode == "enforce":
        assert foreign["counterparty_master_ref"] == {"status": "unmigrated"}


# ================================================================ export
def test_finance_export_has_only_tenant_a_values(monkeypatch):
    from openpyxl import load_workbook
    from app.services import paid_labor
    db = world()
    for mode in MODES:
        set_mode(monkeypatch, db, mode)
        monkeypatch.setattr(paid_labor, "db", db)
        r = client_for(monkeypatch, reports, db, ADMIN).get(
            "/api/reports/company-finance-export?year=2026&month=9&format=xlsx")
        assert r.status_code == 200, r.text
        wb = load_workbook(io.BytesIO(r.content))
        cells = " ".join(str(c.value) for ws in wb for row in ws.iter_rows() for c in row
                         if c.value is not None)
        assert_no_b(db, cells)
        monkeypatch.undo()


# ================================================================ import projection
@pytest.mark.parametrize("mode", MODES)
def test_client_invoice_import_uses_only_tenant_a_owner(monkeypatch, mode):
    db = world()
    ctx = set_mode(monkeypatch, db, mode)
    r = client_for(monkeypatch, projects, db, ADMIN).post("/api/projects/pr-x/import-client-invoice")
    assert r.status_code == 200, r.text
    # own-x is a B company AND an A client: A gets its own client, never B's company
    assert r.json()["invoice_details"]["company_name"] == "А Клиент ООД"
    assert_no_b(db, r.json())
    theirs = run(db["projects"].find_one({"id": "pr-x", "org_id": ORG_B}))
    assert "invoice_details" not in theirs
    if ctx is not None:
        assert r.json()["owner_master_ref"]["status"] in ("mapped", "pending", "unmigrated")
        assert r.json()["owner_master_ref"].get("canonical_id") not in b_canonical_ids(db)


# ================================================================ advances (recipient read + write)
@pytest.mark.parametrize("mode", MODES)
def test_advance_list_never_shows_another_tenants_user(monkeypatch, mode):
    db = world()
    set_mode(monkeypatch, db, mode)
    rows = client_for(monkeypatch, hr, db, ADMIN).get("/api/advances").json()
    row = [a for a in rows if a["id"] == "adv-foreign-user"][0]
    assert row["user_name"] == "Unknown"
    assert_no_b(db, rows)


def test_a_recipient_that_exists_only_in_b_is_refused_without_b_evidence(monkeypatch):
    db = world()
    ctx = enforce_with(monkeypatch, db)
    u1 = run(la.resolve_legacy(ctx, collection="users", legacy_id="u1", mode="enforce",
                               repository=repo(db)))["canonical_id"]
    for body, code in (({"type": "Advance", "user_id": "u-only-b", "person_id": u1},
                        la.REASON_ADVANCE_UNMAPPED),
                       ({"type": "Advance", "user_id": "u3", "person_id": u1},
                        la.REASON_ADVANCE_MISMATCH)):
        before = run(db["advances"].count_documents({})), run(db["finance_payments"].count_documents({}))
        r = client_for(monkeypatch, hr, db, ADMIN).post("/api/advances", json=dict(body, amount=5))
        assert r.status_code == 422 and r.json()["detail"]["error_code"] == code
        assert (run(db["advances"].count_documents({})),
                run(db["finance_payments"].count_documents({}))) == before
        assert_no_b(db, r.json())
    audit_a = run(find_all(db, AUDIT_COLLECTION, {"tenant_id": T_A}))
    assert_no_b(db, audit_a)
    assert run(db[AUDIT_COLLECTION].count_documents({"tenant_id": T_B,
                                                     "action": "master_data.advance.create_refused"})) == 0


# ================================================================ adapter helpers
def test_forged_b_legacy_ref_and_b_rows_never_resolve_in_a(monkeypatch):
    db = world()
    ctx = enforce_with(monkeypatch, db)
    # a Master planted in tenant A that claims B's legacy counterparty
    forged = build_entity(tenant_id=T_A, entity_type="organization", display_name="Подправен",
                          entity_id="org-forged", legacy_refs=[new_legacy_ref(
                              "counterparties", "cp-only-b", ORG_B)])
    run(db["md_organization"].insert_one(forged))
    assert run(la.resolve_legacy(ctx, collection="counterparties", legacy_id="cp-only-b",
                                 mode="enforce", repository=repo(db))) is None
    with pytest.raises(la.LegacyOrgMismatch):
        run(la.resolve_legacy(ctx, collection="counterparties", legacy_id="cp1", org_id=ORG_B,
                              mode="enforce", repository=repo(db)))
    # B's own reverse row for the shared id is a different _id: A's lookup never sees it
    b_row = run(db[lp.REFS_COLLECTION].find_one({"_id": lp.ref_row_id(T_B, "counterparties", "cp1")}))
    a = run(la.resolve_legacy(ctx, collection="counterparties", legacy_id="cp1", mode="enforce",
                              repository=repo(db)))
    assert b_row["entity_id"] != a["canonical_id"] and a["entity"]["tenant_id"] == T_A
    rows = [{"supplier_id": "cp-only-b"}, {"supplier_id": "cp1"}]
    run(la.annotate_refs(ADMIN, rows, {"supplier_id": "counterparties"}))
    assert rows[0]["supplier_master_ref"] == {"status": "unmigrated"}
    assert "org-forged" not in json.dumps(rows)
    assert_no_b(db, rows)
    usage = run(la.count_usage(db, ORG_A, "counterparties", {"id": "cp1"}))
    a_usage = run(db["invoices"].count_documents({"org_id": ORG_A, "supplier_counterparty_id": "cp1"}))
    b_usage = run(db["invoices"].count_documents({"org_id": ORG_B, "supplier_counterparty_id": "cp1"}))
    assert b_usage >= 1 and usage["invoices.supplier_counterparty_id"] == a_usage   # A's own only
    report = run(la.advance_mapping_report(db, tenant_id=T_A, org_id=ORG_A))
    assert_no_b(db, report)


def test_the_dry_run_and_reconciliation_of_a_disclose_nothing_of_b(monkeypatch):
    db = world()
    plan = run(dry(db))
    assert all("other_org" not in inv for inv in plan["inventory"].values())
    assert_no_b(db, plan)
    from app.master_data import legacy_migration as lm
    rec = run(lm.reconcile(Ctx(db), mode="enforce", repository=repo(db)))
    assert_no_b(db, rec)
    mappings = run(lm.list_mappings(Ctx(db), mode="enforce", repository=repo(db), limit=200))
    assert all(m["tenant_id"] == T_A and m["org_id"] == ORG_A for m in mappings)
    assert_no_b(db, mappings)


def test_landmines_off_mode_reports_touch_no_master_data(monkeypatch):
    db = world()
    monkeypatch.delenv(ENV_MODE, raising=False)
    touched = landmines(monkeypatch)
    c = client_for(monkeypatch, reports, db, ADMIN)
    for url in ("/api/prices", "/api/reports/turnover-by-counterparty",
                "/api/reports/turnover-by-client", "/api/reports/turnover-by-counterparty/cp1/invoices"):
        assert c.get(url).status_code == 200
    assert touched == []
